from io import BytesIO
from openai import OpenAI
from app.prompts.document_reader_prompt import DocumentReaderPrompt
from langchain_openai import ChatOpenAI
from dotenv import load_dotenv
from docx import Document
from docx.shared import Inches

_ = load_dotenv()


class AgentDocumentGenerator:
    model = None
    prompt = None

    def __init__(self, model_name="gpt-5.2"):
        self.model = ChatOpenAI(model=model_name)
        self.prompt = DocumentReaderPrompt()

    def generate(self, file_id):
        result = self.model.invoke(
            [
                {"role": "system", "content": self.prompt.prompt_template()},
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": "Analise o documento e gere um resumo objetivo.",
                        },
                        {"type": "file", "file_id": file_id},
                    ],
                },
            ]
        )
        return result.content

template_text = """
"""

document = Document("modelo_acidente_trajeto.docx")
document.add_heading("Modelo de Documento", 0)


# for table in document.tables:
#     for row in table.rows:
#         for cell in row.cells:
#             for paragraph in cell.paragraphs:
#                 print(paragraph.text)

for table_idx, table in enumerate(document.tables):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                print(paragraph.text)
    
    # Adiciona nova linha na tabela com dados do sistema
    new_row = table.add_row()
    
    # Simular dados de um caso real do sistema
    # Em produção, esses dados viriam do banco de dados
    mock_case_data = {
        'id': 1,
        'benefit_number': '123456789',
        'insured_name': 'João da Silva Santos',
        'insured_nit': '123.45678.90-1',
        'accident_date': '15/03/2024',
        'benefit_type': 'B91',
        'status': 'Ativo',
        'value': 'R$ 1.500,00',
        'company': 'Empresa XYZ Ltda',
        'fap_reason': 'Trajeto de Ida ao Trabalho'
    }
    
    for i, cell in enumerate(new_row.cells):
        if i == 0:
            # Primeira coluna: ID
            cell.text = str(mock_case_data['id'])
        elif i == 1:
            cell.text = mock_case_data['benefit_number']
        elif i == 2:
            cell.text = mock_case_data['insured_name']
        elif i == 3:
            cell.text = mock_case_data['insured_nit']
        elif i == 4:
            cell.text = mock_case_data['accident_date']
        elif i == 5:
            cell.text = mock_case_data['benefit_type']
        else:
            # Colunas adicionais
            cell.text = mock_case_data.get('status', f"Dado coluna {i+1}")

for p in document.paragraphs:
    if p.text.startswith("{{") and p.text.endswith("}}"):
        print("\n", p.text)
    # for run in p.runs:
    #     print(run.text)
# agent = AgentDocumentGenerator()
# response = agent.generate()

document.save("modelo_acidente_trajeto_edit.docx")