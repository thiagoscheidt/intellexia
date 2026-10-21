"""Planilha dos benefícios do cálculo FAP — o "Exportar Excel" da tela /fap-panel/calculo/.

Segue os filtros da tela (mesma consulta, mesma ordem, mesma montagem de linha
via ``fap_calculo_query_service.todos``) e o visual da planilha de contestações
do Painel FAP: título, bloco de filtros, totais, cabeçalho azul, Sim/Não e
Conta/Excluído coloridos.

Escrita em modo ``write_only``: o cálculo de um escritório passa de dezenas de
milhares de itens, e o modo normal guardaria a planilha inteira em memória no
worker. Por isso nada de célula mesclada — o título fica só na A1.
"""

from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.cnpj import formatar_cnpj, formatar_cnpj_raiz, formatar_cpf

COLUNAS = [
    # (cabeçalho, largura)
    ('Vigência', 10),
    ('CNPJ estabelecimento', 21),
    ('Empresa', 34),
    ('Tipo', 7),
    ('NB / nº CAT', 15),
    ('Segurado', 32),
    ('NIT', 14),
    ('CPF', 16),
    ('Nascimento', 12),
    ('RMI', 13),
    ('Total pago', 14),
    ('Início', 12),
    ('Cessação', 12),
    ('Despacho', 12),
    ('Início contabilização', 13),
    ('Fim contabilização', 13),
    ('Meses contabilizados', 12),
    ('Data do acidente', 12),
    ('No cálculo atual', 14),
    ('Ausente no cálculo atual', 13),
    ('Contestado', 12),
    ('Cálculo atual', 22),
    ('Protocolo', 22),
]

_AZUL_ESCURO = '0D47A1'
_AZUL = '1976D2'
_FILTRO = 'E8EAF6'

_titulo = (Font(name='Calibri', size=14, bold=True, color='FFFFFF'), PatternFill('solid', fgColor=_AZUL_ESCURO))
_rotulo = (Font(name='Calibri', size=10, bold=True, color='1A237E'), PatternFill('solid', fgColor=_FILTRO))
_valor = (Font(name='Calibri', size=10, color='1A237E'), PatternFill('solid', fgColor=_FILTRO))
_cabecalho = (Font(name='Calibri', size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor=_AZUL))
_sim = (Font(name='Calibri', size=10, bold=True, color='1B5E20'), PatternFill('solid', fgColor='C8E6C9'))
_nao = (Font(name='Calibri', size=10, bold=True, color='B71C1C'), PatternFill('solid', fgColor='FFCDD2'))
_excluido = (Font(name='Calibri', size=10, bold=True, color='1B5E20'), PatternFill('solid', fgColor='C8E6C9'))
_conta = (Font(name='Calibri', size=10, color='37474F'), PatternFill('solid', fgColor='ECEFF1'))
_ausente = (Font(name='Calibri', size=10, bold=True, color='E65100'), PatternFill('solid', fgColor='FFE0B2'))
_fonte_dado = Font(name='Calibri', size=10)

_DATA = 'DD/MM/YYYY'
_DINHEIRO = '"R$" #,##0.00'


def _celula(ws, valor, estilo=None, formato=None, alinhamento=None):
    c = WriteOnlyCell(ws, value=valor)
    if isinstance(valor, str) and valor.startswith('='):
        # O openpyxl grava texto iniciado por "=" como fórmula; a busca da URL
        # (?q=) chega aqui, e "=HYPERLINK(...)" rodaria em quem abrisse o arquivo.
        c.data_type = 's'
    if estilo:
        c.font, c.fill = estilo
    else:
        c.font = _fonte_dado
    if formato:
        c.number_format = formato
    if alinhamento:
        c.alignment = alinhamento
    return c


def _dinheiro(v):
    return float(v) if v is not None else None


def rotulos_dos_filtros(law_firm_id: int, f: dict) -> list[tuple[str, str]]:
    """Filtros aplicados, em texto, para o cabeçalho da planilha."""
    from app.models import FapCompany
    from app.services import fap_group_service

    empresa = 'Todas'
    if f.get('raiz'):
        nome = (FapCompany.query.with_entities(FapCompany.nome)
                .filter_by(law_firm_id=law_firm_id, cnpj=f['raiz']).scalar())
        empresa = f"{formatar_cnpj_raiz(f['raiz'])}{' — ' + nome if nome else ''}"
    grupo = 'Todos'
    if f.get('grupo'):
        grupo = next((o['label'] for o in fap_group_service.select_options(law_firm_id)
                      if o['value'] == f['grupo']), f['grupo'])
    return [
        ('Vigência', str(f['vigencia']) if f.get('vigencia') else 'Todas'),
        ('Grupo', grupo),
        ('Empresa', empresa),
        ('Estabelecimento', formatar_cnpj(f['cnpj']) if f.get('cnpj') else 'Todos'),
        ('Busca (CPF, NIT, NB, CAT)', f.get('q') or '—'),
        ('Tipo', ', '.join(f['tipos']) if f.get('tipos') else 'Todos'),
        ('No cálculo atual', {'conta': 'Conta', 'excluido': 'Excluído'}.get(f.get('situacao'), 'Todos')),
        ('Contestado', {'sim': 'Sim', 'nao': 'Não'}.get(f.get('contestado_arg'), 'Todos')),
    ]


def montar_planilha(itens, rotulos: list[tuple[str, str]], contadores: dict) -> Workbook:
    """``itens``: iterável de dicts de ``fap_calculo_query_service`` (``todos``)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Benefícios do cálculo')
    for i, (_, largura) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    # No write_only, congelar e filtrar se declaram antes da primeira linha.
    linha_cabecalho = 1 + -(-len(rotulos) // 3) + 1 + 1 + 1  # título, filtros, totais, vazia
    ws.freeze_panes = f'A{linha_cabecalho + 1}'
    ws.auto_filter.ref = f'A{linha_cabecalho}:{get_column_letter(len(COLUNAS))}{linha_cabecalho}'

    ws.append([_celula(ws, f'Benefícios do cálculo FAP — exportado em {datetime.now():%d/%m/%Y %H:%M}', _titulo)]
              + [_celula(ws, None, _titulo) for _ in COLUNAS[1:]])

    # Filtros em pares rótulo | valor, três por linha (como a planilha de contestações).
    direita = Alignment(horizontal='right')
    for k in range(0, len(rotulos), 3):
        linha = []
        for rot, val in rotulos[k:k + 3]:
            linha += [_celula(ws, f'{rot}:', _rotulo, alinhamento=direita), _celula(ws, val, _valor)]
        ws.append(linha)

    por_tipo = ', '.join(f'{t}: {n}' for t, n in sorted((contadores.get('por_tipo') or {}).items()) if n)
    totais = [
        ('Itens', f"{contadores.get('total', 0)}" + (f' ({por_tipo})' if por_tipo else '')),
        ('Excluídos do cálculo', str(contadores.get('excluidos', 0))),
        ('Não contestados', str(contadores.get('nao_contestados', 0))),
    ]
    linha = []
    for rot, val in totais:
        linha += [_celula(ws, f'{rot}:', _rotulo, alinhamento=direita), _celula(ws, val, _valor)]
    ws.append(linha)
    ws.append([])

    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.append([_celula(ws, nome, _cabecalho, alinhamento=centro) for nome, _ in COLUNAS])

    for d in itens:
        cat = d.get('tipo') == 'CAT'
        if d.get('ausente_no_calculo_atual'):
            situacao = _celula(ws, 'Ausente', _ausente)
        elif d.get('estado_atual') == 'EXCLUIDO':
            situacao = _celula(ws, 'Excluído', _excluido)
        elif d.get('estado_atual') == 'ORIGINAL':
            situacao = _celula(ws, 'Conta', _conta)
        else:
            situacao = _celula(ws, d.get('estado_atual') or '')
        calculo = ' — '.join(x for x in (
            d.get('calculo_tipo'),
            d['calculo_data'].strftime('%d/%m/%Y') if d.get('calculo_data') else None) if x)
        ws.append([
            _celula(ws, d.get('ano_vigencia')),
            _celula(ws, formatar_cnpj(d.get('cnpj'))),
            _celula(ws, d.get('empresa_nome') or ''),
            _celula(ws, d.get('tipo')),
            _celula(ws, d.get('identificador')),
            _celula(ws, d.get('segurado_nome') or ''),
            _celula(ws, d.get('nit') or ''),
            _celula(ws, formatar_cpf(d.get('cpf')) if d.get('cpf') else ''),
            _celula(ws, d.get('data_nascimento'), formato=_DATA),
            _celula(ws, None if cat else _dinheiro(d.get('valor_mensal')), formato=_DINHEIRO),
            _celula(ws, None if cat else _dinheiro(d.get('total_pago')), formato=_DINHEIRO),
            _celula(ws, d.get('data_inicio'), formato=_DATA),
            _celula(ws, d.get('data_cessacao'), formato=_DATA),
            _celula(ws, d.get('data_despacho'), formato=_DATA),
            _celula(ws, d.get('data_inicio_contabilizacao'), formato=_DATA),
            _celula(ws, d.get('data_fim_contabilizacao'), formato=_DATA),
            _celula(ws, _dinheiro(d.get('meses_contabilizacao')), formato='0.000'),
            _celula(ws, d.get('data_acidente'), formato=_DATA),
            situacao,
            _celula(ws, 'Sim', _ausente) if d.get('ausente_no_calculo_atual') else _celula(ws, ''),
            _celula(ws, 'Sim', _sim) if d.get('contestado') else _celula(ws, 'Não', _nao),
            _celula(ws, calculo),
            _celula(ws, d.get('protocolo') or ''),
        ])
    return wb
