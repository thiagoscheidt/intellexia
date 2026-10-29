"""Planilha de benefícios do Revisor FAP × petição — fonte única.

A tela do Revisor usa este módulo direto; a revisão feita no Claude o recebe
como script (fap_review_claude_service.script_planilha) e roda no ambiente do
Claude, onde estão a planilha e a petição. Por isso o módulo só pode usar a
biblioteca padrão e o openpyxl: nada do app, nada de Flask.

Sem IA: o resultado é o mesmo nos dois caminhos.
"""

import csv
import html
import io
import re
import unicodedata
import zipfile
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def normalizar_cabecalho(valor) -> str:
    """Normaliza cabeçalho de planilha para busca resiliente.

    Além de acento e caixa, derruba pontuação: "Nº", "TESE(S)" e "N° do
    Benefício" viram texto comparável sem uma entrada de alias para cada
    variação de digitação.
    """
    texto = unicodedata.normalize('NFKD', str(valor or ''))
    texto = texto.encode('ascii', 'ignore').decode('ascii')
    texto = re.sub(r'[^a-z0-9]+', ' ', texto.lower())
    return ' '.join(texto.split())


# Cada escritório nomeia a coluna do seu jeito. A lista está em ordem de
# preferência: a primeira que existir na aba é a usada.
CABECALHOS_BENEFICIO = (
    'numero do beneficio',
    'numero beneficio',
    'n do beneficio',
    'no do beneficio',
    'n beneficio',
    'nb',
    'beneficio',
)
CABECALHOS_TESE = ('tese', 'teses')


def _indice_cabecalho(mapa: dict, nomes: tuple) -> int | None:
    """Índice da primeira coluna cujo cabeçalho casa com um dos nomes aceitos."""
    for nome in nomes:
        if nome in mapa:
            return mapa[nome]
    return None


def _indice_tese(mapa: dict) -> int | None:
    """Índice da coluna da tese, aceitando singular, plural e sufixos.

    Casa "TESE", "TESES", "TESE(S)" e "TESES APLICADAS" — mas pela primeira
    palavra, nunca por substring: as vizinhas ("OBS", "Número da CAT") não
    podem ser confundidas com a coluna da tese.
    """
    exato = _indice_cabecalho(mapa, CABECALHOS_TESE)
    if exato is not None:
        return exato
    for cabecalho, indice in mapa.items():
        if cabecalho.split(' ', 1)[0] in CABECALHOS_TESE:
            return indice
    return None


def formatar_numero_beneficio(valor) -> str:
    """Número de benefício como texto legível (o Excel entrega int ou float)."""
    if valor is None or isinstance(valor, bool):
        return ''
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def normalizar_numero_beneficio(valor) -> str:
    """Só os dígitos, para comparar."""
    return re.sub(r'\D+', '', formatar_numero_beneficio(valor))


def documento_cita_beneficio(texto: str, numero_normalizado: str) -> bool:
    """O documento menciona o benefício, com qualquer pontuação entre os dígitos.

    Aceita 1234567890, 123.456.789-0, 123 456 789 0 etc.; dígito colado antes
    ou depois não vale (seria outro número).
    """
    if not texto or not numero_normalizado:
        return False
    separador = r'[\s\.\-\/_;,:\(\)\[\]]*'
    digitos = separador.join(re.escape(d) for d in numero_normalizado)
    return re.search(rf'(?<!\d){digitos}(?!\d)', texto) is not None


def _abas(caminho: str):
    """(nome da aba, linhas) de cada aba. CSV é uma aba só, com o nome do arquivo."""
    if Path(caminho).suffix.lower() == '.csv':
        bruto = Path(caminho).read_bytes()
        for codificacao in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                texto = bruto.decode(codificacao)
                break
            except UnicodeDecodeError:
                continue
        primeira = texto.split('\n', 1)[0]
        delimitador = ';' if primeira.count(';') >= primeira.count(',') and ';' in primeira else ','
        yield Path(caminho).name, list(csv.reader(io.StringIO(texto), delimiter=delimitador))
        return
    if not load_workbook:
        raise ImportError('openpyxl não está instalado: para ler .xlsx ele é necessário')
    livro = load_workbook(filename=caminho, read_only=True, data_only=True)
    try:
        for aba in livro.worksheets:
            yield aba.title, list(aba.iter_rows(values_only=True))
    finally:
        livro.close()


def ler_planilha(caminho: str) -> list[dict]:
    """Linhas com benefício e tese de todas as abas (.xlsx) ou do .csv.

    Planilhas reais costumam ter uma aba por vigência (2021, 2022, ...); abas
    sem as colunas esperadas (ex.: anotações) são ignoradas.
    """
    linhas = []
    abas_com_colunas = 0
    recusadas = []
    for nome_aba, valores in _abas(caminho):
        cabecalho = valores[0] if valores else None
        if not cabecalho:
            recusadas.append(f'{nome_aba}: aba vazia')
            continue
        # Cabeçalho repetido (a planilha real tem "OBS" três vezes): a
        # primeira ocorrência vence, não a última.
        mapa = {}
        for indice, valor in enumerate(cabecalho):
            normalizado = normalizar_cabecalho(valor)
            if normalizado:
                mapa.setdefault(normalizado, indice)
        i_beneficio = _indice_cabecalho(mapa, CABECALHOS_BENEFICIO)
        i_tese = _indice_tese(mapa)
        if i_beneficio is None or i_tese is None:
            faltam = []
            if i_beneficio is None:
                faltam.append('o número do benefício')
            if i_tese is None:
                faltam.append('a tese')
            recusadas.append(f'{nome_aba}: falta {" e ".join(faltam)}')
            continue
        abas_com_colunas += 1
        for linha in valores[1:]:
            tese = ' '.join(str(linha[i_tese] or '').strip().split()) if i_tese < len(linha) else ''
            if not tese:
                continue
            numero = formatar_numero_beneficio(linha[i_beneficio] if i_beneficio < len(linha) else '')
            normalizado = normalizar_numero_beneficio(numero)
            if not normalizado:
                continue
            linhas.append({
                'benefit_number': numero,
                'benefit_number_normalized': normalizado,
                'thesis': tese,
                'sheet_name': nome_aba,
            })
    if not abas_com_colunas:
        # Diz o que faltou em cada aba: sem isso, "cabeçalho com outro nome"
        # e "cabeçalho fora da primeira linha" dão a mesma mensagem.
        detalhe = '; '.join(recusadas[:8])
        if len(recusadas) > 8:
            detalhe += f'; e mais {len(recusadas) - 8} aba(s)'
        raise ValueError(
            'Nenhuma aba da planilha tem as duas colunas obrigatórias na primeira '
            'linha: o número do benefício ("Número do Benefício") e a tese '
            f'("TESE" ou "TESES"). Abas lidas — {detalhe}.'
        )
    return linhas


def conferir_planilha(caminho: str, nome_arquivo: str, texto_documento: str, rotulo_documento: str) -> dict:
    """Cada benefício da planilha (com tese) foi citado no documento?"""
    itens = []
    for linha in ler_planilha(caminho):
        itens.append({
            'benefit_number': linha['benefit_number'],
            'thesis': linha['thesis'],
            'sheet_name': linha.get('sheet_name', ''),
            'found_in_document': documento_cita_beneficio(texto_documento, linha['benefit_number_normalized']),
        })
    citados = sum(1 for i in itens if i['found_in_document'])
    return {
        'source_filename': nome_arquivo,
        'checked_document_label': rotulo_documento,
        'total_rows_with_theses': len(itens),
        'found_count': citados,
        'missing_count': len(itens) - citados,
        'items': itens,
    }


def texto_do_documento(caminho: str) -> str:
    """Texto da petição no ambiente do Claude: .txt, .docx (sem dependência) ou .pdf.

    No .docx, o número pode vir partido entre "runs" do Word; juntar o texto
    das tags sem separador o remonta, e o fim de parágrafo/célula separa.
    """
    extensao = Path(caminho).suffix.lower()
    if extensao in ('.txt', '.md'):
        return Path(caminho).read_text(encoding='utf-8', errors='replace')
    if extensao == '.docx':
        with zipfile.ZipFile(caminho) as pacote:
            xml = pacote.read('word/document.xml').decode('utf-8', errors='replace')
        xml = re.sub(r'</w:p>|<w:br/>|<w:cr/>', '\n', xml)
        xml = re.sub(r'</w:tc>|<w:tab/>', '\t', xml)
        return html.unescape(re.sub(r'<[^>]+>', '', xml))
    if extensao == '.pdf':
        try:
            import pdfplumber
            with pdfplumber.open(caminho) as pdf:
                return '\n'.join(p.extract_text() or '' for p in pdf.pages)
        except ImportError:
            pass
        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                from PyPDF2 import PdfReader
            except ImportError:
                raise ValueError('Sem leitor de PDF aqui: converta a petição para .txt e rode de novo.')
        return '\n'.join(p.extract_text() or '' for p in PdfReader(caminho).pages)
    raise ValueError(f'Formato da petição não suportado ({extensao}): use .docx, .pdf ou .txt.')
