"""
Tools: Base de Jurisprudência
=============================
Decisões FAP (sentença, acórdão, embargos) do escritório, classificadas por
tese e resultado. Só leitura, e só a base de jurisprudência — nada de outros
módulos (processos, peças-modelo).

Tudo passa pelos serviços da tela, para o agente ver exatamente o que a pessoa
vê: a mesma busca (todos os termos, com sinônimos; número de processo com ou
sem pontuação), os mesmos filtros e a mesma citação pronta.

Duas teses por decisão, como na tela:
  - **tese original**: como veio escrita na decisão (N por decisão); o filtro
    junta só grafias que diferem em acento ou maiúsculas;
  - **tese do catálogo**: a padronizada do escritório, pela correspondência.
"""
from __future__ import annotations

import html
import re
from collections import Counter
from datetime import date

from fastmcp.exceptions import ToolError

from mcp_server.tools.pagination import clamp_limit, clamp_offset, page_envelope

MODOS = ('campos', 'inteiro_teor')
ORDENS = ('recentes', 'instancia', 'favoraveis')


def _iso(value):
    return value.isoformat() if value else None


def _texto_puro(trecho_html: str | None) -> str | None:
    """Trecho grifado da tela (<mark>, escapado) → texto com «» no achado."""
    if not trecho_html:
        return None
    texto = trecho_html.replace('<mark>', '«').replace('</mark>', '»')
    return html.unescape(re.sub(r'<[^>]+>', '', texto))


def _curto(texto, tamanho=500):
    texto = (texto or '').strip()
    return texto if len(texto) <= tamanho else texto[:tamanho].rsplit(' ', 1)[0] + '…'


def _url(app_public_url: str | None, caminho: str) -> str | None:
    return f"{app_public_url.rstrip('/')}{caminho}" if app_public_url else None


def _url_decisao(app_public_url, decision_id):
    return _url(app_public_url, f'/process-panel/jurisprudencia/decisao/{decision_id}')


def _catalogo_da_decisao(decisao) -> list[str]:
    return sorted({c.name for t in decisao.theses for c in t.catalog_theses})


def _item(decisao, app_public_url, trecho_html=None) -> dict:
    from app.services import jurisprudence_normalizer as norm

    return {
        'id': decisao.id,
        'processo': decisao.processo,
        'parte_autora': decisao.parte_autora,
        'instancia': norm.TIPO_LABELS.get(decisao.tipo_documento or '', decisao.tipo_documento),
        'tribunal': decisao.tribunal,
        'orgao_julgador': decisao.orgao_julgador,
        'uf': decisao.uf,
        'data_julgamento': _iso(decisao.data_julgamento),
        'resultado': norm.RESULTADO_LABELS.get(decisao.resultado or '', 'não informado'),
        'vigencia_fap': decisao.vigencia_texto,
        'teses_originais': list(decisao.teses_brutas_json or []),
        'teses_catalogo': _catalogo_da_decisao(decisao),
        'motivo_resultado': _curto(decisao.motivo_resultado, 400),
        'trecho': _texto_puro(trecho_html),
        'citacao': norm.citacao(decisao),
        'url': _url_decisao(app_public_url, decisao.id),
    }


def _decisoes_por_id(law_firm_id: int, ids) -> dict:
    from app.models import JurisprudenceDecision

    ids = list(ids)
    if not ids:
        return {}
    return {d.id: d for d in JurisprudenceDecision.query.filter(
        JurisprudenceDecision.law_firm_id == law_firm_id,
        JurisprudenceDecision.id.in_(ids)).all()}


def _decisao(decision_id: int, law_firm_id: int):
    from app.models import JurisprudenceDecision

    decisao = JurisprudenceDecision.query.filter_by(id=decision_id, law_firm_id=law_firm_id).first()
    if decisao is None:
        raise ToolError(f'Decisão {decision_id} não encontrada na base de jurisprudência do escritório.')
    return decisao


# ── Filtros (texto livre do agente → valores da tela) ────────────────

def _lista(valor) -> list[str]:
    if valor is None:
        return []
    if isinstance(valor, str):
        valor = [v for v in re.split(r'[;\n]', valor)]
    return [str(v).strip() for v in valor if str(v).strip()]


def _data(valor, campo):
    if not valor:
        return None
    try:
        return date.fromisoformat(str(valor)[:10])
    except ValueError:
        raise ToolError(f"{campo} inválida: use o formato YYYY-MM-DD (ex.: 2025-01-31).")


def _resolver_resultado(valor: str) -> str:
    from app.services import jurisprudence_normalizer as norm

    k = norm.chave(valor)
    if k.startswith('PARC'):
        return norm.RESULTADO_PARCIAL
    if k.startswith('DESF') or k.startswith('CONTRA') or k.startswith('IMPROC'):
        return norm.RESULTADO_DESFAVORAVEL
    if k.startswith('FAV') or k.startswith('PROC'):
        return norm.RESULTADO_FAVORAVEL
    raise ToolError(f"Resultado '{valor}' não reconhecido. Use: favoravel, parcial ou desfavoravel.")


def _resolver_instancia(valor: str) -> str:
    from app.services import jurisprudence_normalizer as norm

    tipo = norm.tipo_documento(valor)
    if not tipo or (tipo == norm.TIPO_OUTRA and not norm.chave(valor).startswith('OUTR')):
        raise ToolError(f"Instância '{valor}' não reconhecida. Use: sentenca, acordao, embargos ou outra.")
    return tipo


def _resolver_catalogo(law_firm_id: int, nomes: list[str]) -> list[int]:
    """Nome da tese do catálogo (sem diferenciar acento/caixa) ou id → id."""
    from app.models import JudicialLegalThesis
    from app.services import jurisprudence_normalizer as norm

    if not nomes:
        return []
    catalogo = JudicialLegalThesis.query.filter_by(law_firm_id=law_firm_id).all()
    por_chave = {norm.chave(t.name): t.id for t in catalogo}
    por_chave.update({norm.chave((t.key or '').replace('_', ' ')): t.id for t in catalogo if t.key})
    ids_validos = {t.id for t in catalogo}
    saida, desconhecidas = [], []
    for nome in nomes:
        if nome.isdigit() and int(nome) in ids_validos:
            saida.append(int(nome))
        elif norm.chave(nome) in por_chave:
            saida.append(por_chave[norm.chave(nome)])
        else:
            desconhecidas.append(nome)
    if desconhecidas:
        raise ToolError(
            f"Tese do catálogo não encontrada: {', '.join(desconhecidas)}. "
            "Consulte valores_de_filtro_jurisprudencia para os nomes exatos.")
    return saida


def montar_filtros(law_firm_id: int, *, termo=None, teses_catalogo=None, teses_originais=None,
                   resultados=None, tribunais=None, instancias=None, uf=None, vigencia=None,
                   julgada_de=None, julgada_ate=None, ordem='recentes'):
    from app.services import jurisprudence_normalizer as norm
    from app.services.jurisprudence_search_service import Filtros

    if ordem and ordem not in ORDENS:
        raise ToolError(f"Ordem '{ordem}' inválida. Use: {', '.join(ORDENS)}.")
    try:
        vigencia = int(vigencia) if vigencia not in (None, '') else None
    except (TypeError, ValueError):
        raise ToolError('vigencia deve ser um ano (ex.: 2019).')
    return Filtros(
        q=(termo or '').strip()[:300],
        catalogo=_resolver_catalogo(law_firm_id, _lista(teses_catalogo)),
        originais=[norm.chave(o) for o in _lista(teses_originais) if norm.chave(o)],
        resultados=[_resolver_resultado(r) for r in _lista(resultados)],
        tribunais=[t.upper().replace(' ', '') for t in _lista(tribunais)],
        tipos=[_resolver_instancia(i) for i in _lista(instancias)],
        uf=(uf or '').strip().upper()[:2],
        vigencia=vigencia,
        de=_data(julgada_de, 'julgada_de'),
        ate=_data(julgada_ate, 'julgada_ate'),
        ordem=ordem or 'recentes',
        agrupar='decisao',
    )


def _como_tem_decidido(exito: list[dict]) -> dict:
    return {e['rotulo']: {'decisoes': e['n'], 'percentual': round(e['pct'], 1)} for e in exito}


# ── 1. pesquisar_jurisprudencia ───────────────────────────────────────

def search_jurisprudence_handler(law_firm_id: int, app_public_url: str | None = None, *,
                                 modo: str = 'campos', limite: int = 20, deslocamento: int = 0,
                                 **filtros) -> dict:
    from app.services import jurisprudence_search_service as busca

    if modo not in MODOS:
        raise ToolError(f"Modo '{modo}' inválido. Use 'campos' ou 'inteiro_teor'.")
    limite = clamp_limit(limite, 20)
    deslocamento = clamp_offset(deslocamento)
    f = montar_filtros(law_firm_id, **filtros)

    if modo == 'inteiro_teor':
        return _search_full_text(law_firm_id, app_public_url, f, limite, deslocamento)

    r = busca.decisoes_filtradas(law_firm_id, f)
    janela = r['ids'][deslocamento:deslocamento + limite]
    decisoes = _decisoes_por_id(law_firm_id, janela)
    itens = [_item(decisoes[i], app_public_url, busca.trecho_da_decisao(decisoes[i], r['termos']) if f.q else None)
             for i in janela if i in decisoes]
    envelope = page_envelope(r['total'], deslocamento, itens)
    envelope['processos_distintos'] = r['processos']
    envelope['como_tem_decidido'] = _como_tem_decidido(r['exito'])
    if envelope.get('tem_mais'):
        envelope['dica'] = (
            f"Resposta parcial: {len(itens)} de {r['total']}. Para a próxima página repita a busca com "
            f"deslocamento={deslocamento + len(itens)}. Para estatística use panorama_jurisprudencia; "
            "para a lista inteira use exportar_jurisprudencia_excel.")
    return envelope


def _search_full_text(law_firm_id, app_public_url, f, limite, deslocamento) -> dict:
    from app.services import jurisprudence_index_service as indice
    from app.services import jurisprudence_upload_service as envios

    if not f.q:
        raise ToolError("O modo 'inteiro_teor' precisa de um termo.")
    ignorados = [nome for nome, valor in (
        ('teses_catalogo', f.catalogo), ('teses_originais', f.originais), ('uf', f.uf),
        ('vigencia', f.vigencia), ('julgada_de', f.de), ('julgada_ate', f.ate)) if valor]
    achados = indice.buscar_inteiro_teor(law_firm_id, f.q, tribunais=f.tribunais, resultados=f.resultados,
                                         tipos=f.tipos, limite=min(deslocamento + limite + 1, 1000))
    if achados is None:
        raise ToolError('A busca no inteiro teor está indisponível agora. Tente o modo "campos".')
    janela = achados[deslocamento:deslocamento + limite]
    decisoes = _decisoes_por_id(law_firm_id, [a['decision_id'] for a in janela])
    itens = []
    for a in janela:
        d = decisoes.get(a['decision_id'])
        if d is None:
            continue
        item = _item(d, app_public_url, a.get('trecho'))
        item['pagina'] = a.get('page')
        itens.append(item)
    cobertura = envios.cobertura(law_firm_id)
    # O índice não dá total exato: pede-se um a mais só para saber se há
    # próxima página. Quando há, o total é "pelo menos".
    envelope = page_envelope(len(achados), deslocamento, itens)
    if len(achados) > deslocamento + limite:
        envelope['total_e_minimo'] = True
    envelope['cobertura_inteiro_teor'] = {
        'decisoes_na_base': cobertura['total'],
        'com_inteiro_teor_indexado': cobertura['indexadas'],
        'aviso': ('Só as decisões com PDF indexado são pesquisadas no inteiro teor; as demais entram '
                  'apenas pelo modo "campos".') if cobertura['indexadas'] < cobertura['total'] else None,
    }
    if ignorados:
        envelope['filtros_ignorados'] = (
            f"No inteiro teor só valem tribunal, resultado e instância; ignorados: {', '.join(ignorados)}.")
    return envelope


# ── 2. detalhar_decisao ───────────────────────────────────────────────

def decision_detail_handler(decision_id: int, law_firm_id: int, app_public_url: str | None = None) -> dict:
    from app.services import jurisprudence_normalizer as norm
    from app.services import jurisprudence_service as svc

    d = _decisao(decision_id, law_firm_id)
    detalhe = _item(d, app_public_url)
    detalhe.update({
        'classe_processual': d.classe_processual,
        'relator': d.relator,
        'motivo_resultado': d.motivo_resultado,
        'resumo': d.resumo_executivo,
        'ementa': d.ementa,
        'ementa_tipo': ('resumo — não citar como transcrição' if d.ementa_modo == 'resumo'
                        else 'transcrição' if d.ementa else None),
        'fundamentos': list(d.fundamentos_json or []),
        'argumentos_acolhidos': list(d.argumentos_acolhidos_json or []),
        'argumentos_rejeitados': list(d.argumentos_rejeitados_json or []),
        'precedentes_citados': [
            {'texto': p['texto'], 'decisao_na_base_id': p['decision_id']}
            for p in svc.precedentes_com_link(d)
        ],
        'pdf_disponivel': bool(d.pdf_path),
        'origem': 'planilha importada' if d.source == 'planilha' else 'PDF lido pela IA',
        'outras_decisoes_do_processo': [
            {'id': x.id, 'instancia': norm.TIPO_LABELS.get(x.tipo_documento or '', x.tipo_documento),
             'data_julgamento': _iso(x.data_julgamento),
             'resultado': norm.RESULTADO_LABELS.get(x.resultado or '', 'não informado')}
            for x in svc.decisoes_do_processo(d) if x.id != d.id
        ],
    })
    return detalhe


# ── 3. panorama_jurisprudencia ────────────────────────────────────────

def jurisprudence_overview_handler(law_firm_id: int, app_public_url: str | None = None, **filtros) -> dict:
    """ "Como tem decidido": resultado por tribunal e por instância, viradas no
    acórdão e as favoráveis mais recentes — sobre o recorte dos filtros."""
    from app.services import jurisprudence_normalizer as norm
    from app.services import jurisprudence_search_service as busca
    from app.services import jurisprudence_service as svc
    from app.models import JurisprudenceDecision

    f = montar_filtros(law_firm_id, **filtros)
    r = busca.decisoes_filtradas(law_firm_id, f)
    decisoes = list(_decisoes_por_id(law_firm_id, r['ids']).values())
    rotulo = norm.RESULTADO_LABELS

    def tabela(chave):
        grupos: dict[str, Counter] = {}
        for d in decisoes:
            grupos.setdefault(chave(d) or 'não informado', Counter())[rotulo.get(d.resultado or '', 'não informado')] += 1
        return [{'grupo': g, 'decisoes': sum(c.values()), **dict(c)}
                for g, c in sorted(grupos.items(), key=lambda kv: -sum(kv[1].values()))]

    # Viradas: processos do recorte cuja sentença perdeu e o acórdão melhorou.
    processos = {d.processo_digits for d in decisoes if d.processo_digits}
    trilhas: dict[str, list] = {}
    if processos:
        for d in JurisprudenceDecision.query.filter(
                JurisprudenceDecision.law_firm_id == law_firm_id,
                JurisprudenceDecision.processo_digits.in_(processos)).all():
            trilhas.setdefault(d.processo_digits, []).append(d)
    viradas = []
    for trilha in trilhas.values():
        trilha = svc.ordenar_trilha(trilha)
        if busca.virou_no_acordao(trilha):
            acordao = next(x for x in reversed(trilha) if x.tipo_documento == norm.TIPO_ACORDAO)
            viradas.append({'processo': acordao.processo, 'parte_autora': acordao.parte_autora,
                            'acordao_id': acordao.id, 'tribunal': acordao.tribunal,
                            'resultado_do_acordao': rotulo.get(acordao.resultado or ''),
                            'citacao': norm.citacao(acordao), 'url': _url_decisao(app_public_url, acordao.id)})

    a_favor = sorted((d for d in decisoes if d.resultado in (norm.RESULTADO_FAVORAVEL, norm.RESULTADO_PARCIAL)),
                     key=lambda d: (d.data_julgamento or date.min, norm.TIPO_ORDEM.get(d.tipo_documento, 0), d.id),
                     reverse=True)[:8]
    return {
        'recorte': {k: v for k, v in filtros.items() if v not in (None, '', [])} or 'toda a base',
        'decisoes': r['total'],
        'processos': r['processos'],
        'como_tem_decidido': _como_tem_decidido(r['exito']),
        'por_tribunal': tabela(lambda d: d.tribunal),
        'por_instancia': tabela(lambda d: norm.TIPO_LABELS.get(d.tipo_documento or '')),
        'viradas_no_acordao': viradas[:10],
        'total_viradas': len(viradas),
        'favoraveis_mais_recentes': [_item(d, app_public_url) for d in a_favor],
        'observacao': ('Tese do catálogo só conta decisões cuja tese já foi ligada ao catálogo na '
                       'correspondência de teses; a tese original conta tudo o que está escrito na decisão.'),
    }


# ── 4. valores_de_filtro_jurisprudencia ───────────────────────────────

def jurisprudence_filter_values_handler(law_firm_id: int) -> dict:
    from app.models import db, JudicialLegalThesis, JurisprudenceDecision as D
    from app.services import jurisprudence_normalizer as norm
    from app.services import jurisprudence_search_service as busca

    r = busca.buscar(law_firm_id, busca.Filtros(), por_pagina=1)
    catalogo = {t.id: t.name for t in JudicialLegalThesis.query.filter_by(law_firm_id=law_firm_id, is_active=True)}
    datas = db.session.query(db.func.min(D.data_julgamento), db.func.max(D.data_julgamento)).filter(
        D.law_firm_id == law_firm_id).first()
    ordem = lambda contagem: sorted(contagem.items(), key=lambda kv: (-kv[1], str(kv[0])))  # noqa: E731
    return {
        'teses_catalogo': [{'nome': catalogo[i], 'decisoes': r['facetas']['catalogo'].get(i, 0)}
                           for i in sorted(catalogo, key=lambda i: (-r['facetas']['catalogo'].get(i, 0), catalogo[i]))],
        'teses_originais': [{'nome': r['nomes_originais'].get(k, k), 'decisoes': n}
                            for k, n in ordem(r['facetas']['originais'])],
        'tribunais': [{'sigla': t, 'decisoes': n} for t, n in ordem(r['facetas']['tribunal'])],
        'instancias': [{'valor': t, 'rotulo': norm.TIPO_LABELS[t], 'decisoes': n}
                       for t, n in ordem(r['facetas']['tipo'])],
        'resultados': [{'valor': k, 'rotulo': norm.RESULTADO_LABELS[k], 'decisoes': n}
                       for k, n in ordem(r['facetas']['resultado'])],
        'ufs': [{'uf': u, 'decisoes': n} for u, n in ordem(r['facetas']['uf'])],
        'anos_de_vigencia': r['anos_vigencia'],
        'periodo_de_julgamento': {'de': _iso(datas[0]) if datas else None, 'ate': _iso(datas[1]) if datas else None},
        'total_decisoes': r['total'],
    }


# ── 6. decisoes_parecidas ─────────────────────────────────────────────

def similar_decisions_handler(decision_id: int, law_firm_id: int, app_public_url: str | None = None,
                              limite: int = 6) -> dict:
    from app.services import jurisprudence_index_service as indice

    d = _decisao(decision_id, law_firm_id)
    limite = max(1, min(int(limite or 6), 20))
    achados = indice.parecidas(d, limite=limite)
    if achados is None:
        raise ToolError('Decisões parecidas indisponíveis para esta decisão: ela ainda não está no índice '
                        'ou o índice não respondeu. Use pesquisar_jurisprudencia com as teses dela.')
    decisoes = _decisoes_por_id(law_firm_id, [a['decision_id'] for a in achados])
    return {
        'decisao_de_referencia': _item(d, app_public_url),
        'parecidas': [{**_item(decisoes[a['decision_id']], app_public_url), 'semelhanca': round(a['score'], 3)}
                      for a in achados if a['decision_id'] in decisoes],
        'criterio': 'semelhança de conteúdo (índice próprio da base de jurisprudência), não palavra nem tese',
    }


# ── 7. exportar_jurisprudencia_excel ──────────────────────────────────

_COLUNAS = [
    ('Processo', 18), ('Parte autora', 32), ('Instância', 14), ('Classe', 22), ('Tribunal', 9),
    ('Órgão julgador', 30), ('UF', 5), ('Relator(a)/Juiz(a)', 28), ('Data do julgamento', 13),
    ('Resultado', 20), ('Por que esse resultado', 60), ('Vigência FAP', 12), ('Teses originais', 45),
    ('Teses do catálogo', 40), ('Fundamentos', 60), ('Argumentos acolhidos', 50),
    ('Argumentos rejeitados', 50), ('Precedentes citados', 45), ('Ementa', 60), ('Ementa é', 14),
    ('Resumo do caso', 60), ('Citação', 50), ('Link', 40),
]


def export_jurisprudence_excel_handler(law_firm_id: int, mcp_public_url: str, app_public_url: str | None = None,
                                       **filtros) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.services import jurisprudence_normalizer as norm
    from app.services import jurisprudence_search_service as busca
    from mcp_server.tools.exports import MAX_EXPORT_ROWS, _download_result, _save_workbook

    f = montar_filtros(law_firm_id, **filtros)
    r = busca.decisoes_filtradas(law_firm_id, f)
    if not r['ids']:
        return {'erro': 'Nenhuma decisão encontrada com esses filtros.', 'total_linhas': 0}
    ids = r['ids'][:MAX_EXPORT_ROWS]
    decisoes = _decisoes_por_id(law_firm_id, ids)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Jurisprudência'
    ws.append([c for c, _ in _COLUNAS])
    for i, (_, largura) in enumerate(_COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        cel = ws.cell(row=1, column=i)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor='4F46E5')
    junta = '; '.join
    for decision_id in ids:
        d = decisoes.get(decision_id)
        if d is None:
            continue
        ws.append([
            d.processo, d.parte_autora, norm.TIPO_LABELS.get(d.tipo_documento or '', d.tipo_documento),
            d.classe_processual, d.tribunal, d.orgao_julgador, d.uf, d.relator, d.data_julgamento,
            norm.RESULTADO_LABELS.get(d.resultado or '', ''), d.motivo_resultado, d.vigencia_texto,
            junta(d.teses_brutas_json or []), junta(_catalogo_da_decisao(d)), junta(d.fundamentos_json or []),
            junta(d.argumentos_acolhidos_json or []), junta(d.argumentos_rejeitados_json or []),
            junta(d.precedentes_json or []), d.ementa,
            ('resumo' if d.ementa_modo == 'resumo' else 'transcrição') if d.ementa else '',
            d.resumo_executivo, norm.citacao(d), _url_decisao(app_public_url, d.id),
        ])
    for linha in ws.iter_rows(min_row=2):
        linha[8].number_format = 'DD/MM/YYYY'
        for cel in linha:
            cel.alignment = Alignment(vertical='top', wrap_text=False)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions

    rel_path = _save_workbook(wb, law_firm_id, 'jurisprudencia')
    resultado = _download_result(rel_path, mcp_public_url, len(ids))
    resultado['como_tem_decidido'] = _como_tem_decidido(r['exito'])
    if r['total'] > MAX_EXPORT_ROWS:
        resultado['aviso'] = f'Resultado truncado em {MAX_EXPORT_ROWS} linhas (total: {r["total"]}).'
    return resultado
