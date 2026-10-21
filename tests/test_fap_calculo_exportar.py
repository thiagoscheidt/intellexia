#!/usr/bin/env python3
"""
"Exportar Excel" da tela de Benefícios do cálculo FAP (/fap-panel/calculo/).

A planilha segue os filtros da tela: mesma consulta, mesma ordem e a mesma
montagem de linha (nome revelado, contestado, situação no cálculo). Escritório
descartável, removido no fim; CPFs fictícios.

    uv run python tests/test_fap_calculo_exportar.py
"""

import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from main import app
from app.models import (
    db, FapCalculoEstabelecimento, FapCalculoItem, FapCompany, LawFirm, SeguradoNome,
    User, UserPageVisit,
)

CNPJ_FIRM = '00000000000999'
_falhas = []


def check(nome: str, condicao: bool, detalhe: str = '') -> None:
    if condicao:
        print(f'  ✅ {nome}')
    else:
        print(f'  ❌ {nome}{" — " + detalhe if detalhe else ""}')
        _falhas.append(nome)


def _setup():
    f = LawFirm.query.filter_by(cnpj=CNPJ_FIRM).first()
    if f:
        _cleanup(f)
    firm = LawFirm(name='Escritório Teste Exportar Cálculo', cnpj=CNPJ_FIRM)
    db.session.add(firm)
    db.session.flush()
    db.session.add(FapCompany(law_firm_id=firm.id, cnpj='07718633', nome='EMPRESA TESTE SA', synced_at=datetime.utcnow()))
    itens = [
        ('B91', '9933000001', '12345678909', 'ORIGINAL', 2026, 1500.50),
        ('B91', '9933000002', '11144477735', 'EXCLUIDO', 2026, 800.00),
        ('B94', '9933000003', None, 'ORIGINAL', 2026, 300.00),
        ('CAT', '9933000004', None, 'ORIGINAL', 2026, None),
        ('B91', '9933000005', None, 'ORIGINAL', 2025, 100.00),
    ]
    for tipo, ident, cpf, estado, ano, pago in itens:
        db.session.add(FapCalculoItem(
            law_firm_id=firm.id, ano_vigencia=ano, cnpj='07718633008163', cnpj_raiz='07718633',
            tipo=tipo, identificador=ident, cpf=cpf, estado_atual=estado,
            data_nascimento=date(1990, 1, 1) if cpf else None, total_pago=pago,
            data_acidente=date(2024, 3, 1) if tipo == 'CAT' else None))
    db.session.add(SeguradoNome(law_firm_id=firm.id, cpf='12345678909', nome='MARIA DE TESTE',
                                status='encontrado', nascimento_diverge=False, fonte='teste',
                                consultado_em=datetime.utcnow()))
    user = User(law_firm_id=firm.id, name='t', email='user__test_exp_calc@example.com', role='admin', is_active=True)
    user.set_password('x')
    db.session.add(user)
    db.session.commit()
    return firm, user


def _cleanup(firm):
    for u in User.query.filter_by(law_firm_id=firm.id).all():
        UserPageVisit.query.filter_by(user_id=u.id).delete()
        db.session.delete(u)
    SeguradoNome.query.filter_by(law_firm_id=firm.id).delete()
    FapCalculoItem.query.filter_by(law_firm_id=firm.id).delete()
    FapCalculoEstabelecimento.query.filter_by(law_firm_id=firm.id).delete()
    FapCompany.query.filter_by(law_firm_id=firm.id).delete()
    db.session.delete(firm)
    db.session.commit()


def _planilha(client, query=''):
    r = client.get('/fap-panel/calculo/exportar.xlsx' + query)
    assert r.status_code == 200, r.status_code
    ws = load_workbook(BytesIO(r.data)).active
    linhas = [[c.value for c in row] for row in ws.iter_rows()]
    cab = next(i for i, l in enumerate(linhas) if l and l[0] == 'Vigência')
    return ws, linhas[:cab], linhas[cab], linhas[cab + 1:]


def run(firm, user):
    c = app.test_client()
    with c.session_transaction() as s:
        s.update(user_id=user.id, law_firm_id=firm.id, user_role='admin',
                 user_module_permissions=user.get_module_permissions())

    print('\n1. Sem filtro: tudo, na ordem da tela')
    ws, topo, cab, dados = _planilha(c)
    check('5 itens', len(dados) == 5, str(len(dados)))
    check('vigência mais nova primeiro', dados[0][0] == 2026 and dados[-1][0] == 2025)
    check('cabeçalho congelado logo abaixo', ws.freeze_panes == f'A{len(topo) + 2}', str(ws.freeze_panes))
    linha = {d[4]: dict(zip(cab, d)) for d in dados}
    l1 = linha['9933000001']
    check('nome revelado na coluna Segurado', l1['Segurado'] == 'MARIA DE TESTE')
    check('CPF formatado', l1['CPF'] == '123.456.789-09')
    check('CNPJ formatado', l1['CNPJ estabelecimento'] == '07.718.633/0081-63')
    check('total pago numérico', l1['Total pago'] == 1500.5)
    check('Conta / Excluído', l1['No cálculo atual'] == 'Conta'
          and linha['9933000002']['No cálculo atual'] == 'Excluído')
    check('Contestado = Não', l1['Contestado'] == 'Não')
    check('CAT sem valor e com data do acidente', linha['9933000004']['Total pago'] is None
          and linha['9933000004']['Data do acidente'] is not None)
    check('empresa pelo nome', l1['Empresa'] == 'EMPRESA TESTE SA')

    print('\n2. Com filtro: segue a tela')
    ws, topo, cab, dados = _planilha(c, '?vigencia=2026&tipo=B91&situacao=conta')
    check('só B91 de 2026 que contam', [d[4] for d in dados] == ['9933000001'], str([d[4] for d in dados]))
    texto_topo = ' '.join(str(v) for l in topo for v in l if v)
    check('filtros escritos no topo', 'Vigência:' in texto_topo and '2026' in texto_topo
          and 'B91' in texto_topo and 'Conta' in texto_topo, texto_topo[:300])
    check('total do topo bate com as linhas', '1 (B91: 1)' in texto_topo, texto_topo[:300])

    print('\n3. Busca com "=" não vira fórmula')
    r = c.get('/fap-panel/calculo/exportar.xlsx?q==HYPERLINK("http://x","clique")')
    ws = load_workbook(BytesIO(r.data)).active
    celulas = [cel for row in ws.iter_rows() for cel in row if isinstance(cel.value, str) and 'HYPERLINK' in cel.value]
    check('valor gravado como texto', celulas and all(cel.data_type == 's' for cel in celulas),
          str([(cel.value, cel.data_type) for cel in celulas]))

    print('\n4. Botão na tela leva os filtros')
    html = c.get('/fap-panel/calculo/?vigencia=2026&tipo=B91').get_data(as_text=True)
    check('link de exportar com os filtros',
          '/fap-panel/calculo/exportar.xlsx?' in html and 'vigencia=2026' in html.split('exportar.xlsx?')[1][:200]
          and 'tipo=B91' in html.split('exportar.xlsx?')[1][:200])


def main() -> int:
    with app.app_context():
        firm, user = _setup()
        try:
            run(firm, user)
        finally:
            db.session.rollback()
            _cleanup(firm)

    print('\n' + '=' * 62)
    if _falhas:
        print(f'❌ {len(_falhas)} verificação(ões) falharam:')
        for nome in _falhas:
            print(f'   - {nome}')
        return 1
    print('✅ Tudo verde')
    return 0


if __name__ == '__main__':
    sys.exit(main())
