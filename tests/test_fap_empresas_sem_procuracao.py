"""
Empresas sem procuração vigente na tela de Procurações FAP.

A tela é uma linha por procuração; a empresa que entra só pelo vínculo gov.br
(grupo Vale) não tem linha nenhuma e sumia dali. Testa a regra pura de
"vigente" e as três saídas — lista geral, filtro por empresa e Excel.

    uv run python tests/test_fap_empresas_sem_procuracao.py
"""

import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from main import app
from app.models import User
from app.services import fap_procuracoes_service as svc

HOJE = date(2026, 10, 13)


def _proc(situacao, inicio, fim, cadastro=None, protocolo='1'):
    return SimpleNamespace(situacao_codigo=situacao, data_inicio=inicio, data_fim=fim,
                           data_cadastro=cadastro, protocolo=protocolo)


def test_vigente_exige_deferida_e_data():
    assert svc.esta_vigente(_proc('DEFERIDA', date(2026, 1, 1), date(2026, 12, 31)), HOJE)
    assert svc.esta_vigente(_proc('DEFERIDA', None, None), HOJE)
    assert svc.esta_vigente(_proc('DEFERIDA', date(2026, 1, 1), HOJE), HOJE)             # vence hoje: ainda vale
    assert not svc.esta_vigente(_proc('DEFERIDA', date(2025, 1, 1), date(2026, 10, 12)), HOJE)   # deferida vencida
    assert not svc.esta_vigente(_proc('DEFERIDA', date(2026, 11, 1), date(2027, 1, 1)), HOJE)    # ainda não começou
    assert not svc.esta_vigente(_proc('EXCLUIDA', date(2026, 1, 1), date(2026, 12, 31)), HOJE)   # no prazo, mas excluída
    assert not svc.esta_vigente(_proc('EXPIRADA', date(2024, 1, 1), date(2024, 12, 31)), HOJE)


def test_ultima_procuracao():
    antiga = _proc('EXPIRADA', date(2023, 1, 1), date(2023, 12, 31), protocolo='A')
    recente = _proc('EXPIRADA', date(2025, 1, 1), date(2025, 12, 31), protocolo='B')
    vigente = _proc('DEFERIDA', date(2026, 1, 1), date(2026, 12, 31), protocolo='C')
    assert svc.ultima_procuracao([antiga, recente], HOJE).protocolo == 'B'
    assert svc.ultima_procuracao([antiga, recente, vigente], HOJE) is None      # tem vigente: fica de fora
    assert svc.ultima_procuracao([], HOJE) is None                            # nunca teve
    mesmo_fim = _proc('EXCLUIDA', date(2025, 1, 1), date(2025, 12, 31), datetime(2025, 6, 1), 'D')
    recente.data_cadastro = datetime(2025, 1, 1)
    assert svc.ultima_procuracao([recente, mesmo_fim], HOJE).protocolo == 'D'  # empate: a cadastrada por último


def _cliente():
    u = User.query.filter_by(role='admin').first()
    c = app.test_client()
    with c.session_transaction() as s:
        s['user_id'] = u.id; s['law_firm_id'] = u.law_firm_id; s['user_role'] = u.role; s['user_name'] = u.name
    return c, u.law_firm_id


def test_telas_e_excel():
    from openpyxl import load_workbook
    with app.app_context():
        c, lf = _cliente()
        todas = svc.empresas_sem_procuracao_vigente(lf)
        nunca = [e for e in todas if e['ultima'] is None]

        # Filtro por empresa: uma linha por empresa, nenhuma de procuração.
        sem = [e for e in todas if not e['por_vinculo']]
        egov = [e for e in todas if e['por_vinculo']]
        html = c.get(f'/fap-panel/procuracoes?situacao={svc.SEM_VIGENTE}').get_data(as_text=True)
        assert html.count('class="linha-sem-vigente"') == len(sem), (html.count('linha-sem-vigente'), len(sem))
        assert 'class="linha-egov"' not in html  # quem tem vínculo e-gov tem acesso: não entra em "sem vigente"
        html = c.get(f'/fap-panel/procuracoes?situacao={svc.VINCULADA_EGOV}').get_data(as_text=True)
        assert html.count('class="linha-egov"') == len(egov) and 'class="linha-sem-vigente"' not in html

        # Lista geral: só quem nunca teve procuração (as outras já estão nas próprias linhas).
        html = c.get('/fap-panel/procuracoes').get_data(as_text=True)
        assert html.count('class="linha-sem-vigente"') + html.count('class="linha-egov"') == len(nunca)
        assert 'Empresas sem procuração vigente' in html

        # Filtro que só existe em procuração tira as empresas da visão.
        html = c.get('/fap-panel/procuracoes?situacao=DEFERIDA').get_data(as_text=True)
        assert 'class="linha-sem-vigente"' not in html

        # Excel acompanha a tela.
        r = c.get(f'/fap-panel/procuracoes/export-excel?situacao={svc.SEM_VIGENTE}')
        assert r.status_code == 200
        ws = load_workbook(BytesIO(r.data)).active
        linhas = [row for row in ws.iter_rows(values_only=True) if row[2] and str(row[2]).startswith(svc.ROTULO_SEM_VIGENTE)]
        assert len(linhas) == len(sem), (len(linhas), len(sem))
        print(f'      ({len(sem)} sem vigente, {len(egov)} pelo e-gov, {len(nunca)} sem histórico)')


def test_linha_de_quem_nunca_teve_procuracao():
    """O banco de dev pode não ter empresa só por vínculo — simula uma."""
    falsa = [{'cnpj_raiz': '33592510', 'nome': 'VALE S.A.', 'por_vinculo': True,
              'vinculo_desde': date(2019, 2, 28), 'ultima': None}]
    original = svc.empresas_sem_procuracao_vigente
    svc.empresas_sem_procuracao_vigente = lambda *a, **k: falsa
    try:
        with app.app_context():
            c, _ = _cliente()
            html = c.get('/fap-panel/procuracoes').get_data(as_text=True)
            assert 'VALE S.A.' in html and 'Nunca teve procuração' in html and 'Vinculada pelo e-gov' in html
            assert 'desde 28/02/2019' in html
            assert '28/02/2019' in html and 'não informada' in html  # vigência do vínculo
            assert 'vinculada(s) pelo e-gov' in html and 'sem procuração vigente</span>' not in html.split('</h6>')[0][-600:]
            filtrado = c.get(f'/fap-panel/procuracoes?situacao={svc.VINCULADA_EGOV}').get_data(as_text=True)
            assert 'VALE S.A.' in filtrado and 'class="linha-egov"' in filtrado
            assert '33.592.510' in html
            r = c.get('/fap-panel/procuracoes/export-excel')
            from openpyxl import load_workbook
            ws = load_workbook(BytesIO(r.data)).active
            linha = next(row for row in ws.iter_rows(values_only=True) if row[3] == 'VALE S.A.')
            assert linha[2] == 'Vinculada pelo e-gov desde 28/02/2019 (nunca teve procuração)', linha
            assert linha[7] == '28/02/2019' and linha[8] == 'não informada', linha
    finally:
        svc.empresas_sem_procuracao_vigente = original


if __name__ == '__main__':
    falhas = 0
    for nome, fn in list(globals().items()):
        if nome.startswith('test_') and callable(fn):
            try:
                fn()
                print(f'OK    {nome}')
            except AssertionError as e:
                falhas += 1
                print(f'FALHA {nome}: {e}')
    sys.exit(1 if falhas else 0)
