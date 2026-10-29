"""
Planilha de benefícios na revisão feita no Claude.

Nem a planilha nem a petição chegam ao IntellexIA: a ferramenta
conferir_planilha_beneficios entrega um script montado do mesmo módulo que a
tela usa (app/utils/planilha_beneficios.py), o Claude roda no ambiente dele e
manda o resultado na conferência. Aqui se confere que:

  1. o script roda isolado (sem o app) e dá o mesmo resultado da tela;
  2. CSV (; ou ,) e DOCX são lidos pelo mesmo caminho;
  3. a conferência guarda a checagem só quando veio do script atual, com as
     contagens recontadas, e ela vai para a tela como na revisão da web.

Uso: uv run python tests/test_fap_review_planilha_claude.py
"""

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402

from main import app  # noqa: E402
from app.models import FapReviewDraft, LawFirm, User, db  # noqa: E402
from app.services import fap_review_claude_service as claude_svc  # noqa: E402
from app.utils import planilha_beneficios as planilha  # noqa: E402

_falhas: list[str] = []


def check(nome: str, condicao: bool, detalhe: str = '') -> None:
    if condicao:
        print(f'  ✅ {nome}')
    else:
        print(f'  ❌ {nome}' + (f' — {detalhe}' if detalhe else ''))
        _falhas.append(nome)


PETICAO = (
    'DOS BENEFÍCIOS\n'
    'O benefício NB 640.713.298-7 foi concedido sem nexo.\n'
    'Também o NB 6222222222, de trajeto.\n'
    'Número 16407132987 não é benefício (11 dígitos).\n'
)


def _xlsx(pasta: Path) -> Path:
    caminho = pasta / 'beneficios.xlsx'
    wb = Workbook()
    aba = wb.active
    aba.title = '2021'
    aba.append(['Nº do Benefício', 'TESE(S)'])
    aba.append([6407132987, 'BENEFÍCIO CANCELADO'])
    aba.append(['1111111111', ''])
    outra = wb.create_sheet('2022')
    outra.append(['NB', 'TESES'])
    outra.append(['622.222.222-2', 'TRAJETO'])
    outra.append(['6333333333', 'B31 INDEVIDO'])
    wb.create_sheet('Notas').append(['texto livre'])
    wb.save(caminho)
    return caminho


def _docx(pasta: Path, texto: str) -> Path:
    """DOCX mínimo, com o número partido em dois 'runs' como o Word faz."""
    caminho = pasta / 'peticao.docx'
    paragrafos = []
    for linha in texto.splitlines():
        if '640.713' in linha:
            corpo = ('<w:r><w:t>O benefício NB 640.71</w:t></w:r>'
                     '<w:r><w:t>3.298-7 foi concedido sem nexo.</w:t></w:r>')
        else:
            corpo = f'<w:r><w:t>{linha}</w:t></w:r>'
        paragrafos.append(f'<w:p>{corpo}</w:p>')
    documento = ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
                 'wordprocessingml/2006/main"><w:body>' + ''.join(paragrafos) + '</w:body></w:document>')
    with zipfile.ZipFile(caminho, 'w') as z:
        z.writestr('[Content_Types].xml', '<Types/>')
        z.writestr('word/document.xml', documento)
    return caminho


def _rodar_script(pasta: Path, planilha_path: Path, peticao_path: Path) -> dict:
    script = pasta / 'conferir_planilha.py'
    script.write_text(claude_svc.script_planilha()['script'], encoding='utf-8')
    # -I e cwd fora do projeto: o script não pode depender de nada do app.
    saida = subprocess.run([sys.executable, '-I', str(script), str(planilha_path), str(peticao_path)],
                           capture_output=True, text=True, cwd=pasta, timeout=60)
    if saida.returncode != 0:
        return {'_erro': saida.stderr[-500:]}
    return json.loads(saida.stdout)


def test_modulo_e_script():
    print('\n1) Script isolado = resultado da tela')
    with tempfile.TemporaryDirectory() as tmp:
        pasta = Path(tmp)
        xlsx = _xlsx(pasta)
        txt = pasta / 'peticao.txt'
        txt.write_text(PETICAO, encoding='utf-8')

        tela = planilha.conferir_planilha(str(xlsx), 'beneficios.xlsx', PETICAO, 'peticao.txt')
        check('módulo: 3 linhas com tese, 2 citadas, 1 faltando',
              (tela['total_rows_with_theses'], tela['found_count'], tela['missing_count']) == (3, 2, 1), str(tela))

        r = _rodar_script(pasta, xlsx, txt)
        check('script roda isolado, com -I', '_erro' not in r, r.get('_erro', ''))
        check('script dá os mesmos itens da tela', r.get('items') == tela['items'], str(r)[:300])
        check('script traz a versão para a conferência',
              r.get('script_versao') == claude_svc.script_planilha()['versao'])
        check('script não traz docstring nem comentário (o Claude copia o texto)',
              '"""' not in claude_svc.script_planilha()['script'])

        docx = _docx(pasta, PETICAO)
        r_docx = _rodar_script(pasta, xlsx, docx)
        check('DOCX: número partido entre runs é achado', r_docx.get('items') == tela['items'], str(r_docx)[:300])

        csv_pv = pasta / 'beneficios.csv'
        csv_pv.write_text('Número do Benefício;Tese\n640.713.298-7;CANCELADO\n6999999999;TRAJETO\n',
                          encoding='utf-8-sig')
        r_csv = planilha.conferir_planilha(str(csv_pv), 'beneficios.csv', PETICAO, 'peticao')
        check('CSV com ponto e vírgula', (r_csv['total_rows_with_theses'], r_csv['found_count']) == (2, 1),
              str(r_csv))
        check('CSV: a aba é o nome do arquivo', r_csv['items'][0]['sheet_name'] == 'beneficios.csv')
        csv_vg = pasta / 'virgula.csv'
        csv_vg.write_text('NB,TESES\n6222222222,TRAJETO\n', encoding='latin-1')
        r_vg = planilha.conferir_planilha(str(csv_vg), 'virgula.csv', PETICAO, 'peticao')
        check('CSV com vírgula e em latin-1', r_vg['found_count'] == 1, str(r_vg))

        sem_coluna = pasta / 'ruim.csv'
        sem_coluna.write_text('Nome;Valor\nx;1\n', encoding='utf-8')
        try:
            planilha.ler_planilha(str(sem_coluna))
            check('planilha sem as colunas: erro que explica', False)
        except ValueError as e:
            check('planilha sem as colunas: erro que explica', 'número do benefício' in str(e), str(e))


def _resultado():
    return {
        'theses': [{'thesis': 'Trajeto', 'benefit_number': '6222222222'}],
        'findings': [{'category': 'FORMAL', 'severity': 'FORMAL', 'description': 'Numeração ausente',
                      'location': 'Fecho', 'location_excerpt': None, 'correction': 'Numerar',
                      'manual_reference': None, 'is_new_pattern': False}],
        'missing_documents': [],
        'executive_summary': {},
    }


def test_conferencia(firm_id, user_id):
    print('\n2) Conferência com a planilha')
    versao = claude_svc.script_planilha()['versao']
    checagem = {
        'source_filename': 'beneficios.xlsx', 'checked_document_label': 'peticao.docx',
        'total_rows_with_theses': 99, 'found_count': 99, 'missing_count': 0,
        'items': [
            {'benefit_number': '6407132987', 'thesis': 'CANCELADO', 'sheet_name': '2021', 'found_in_document': True},
            {'benefit_number': '6333333333', 'thesis': 'B31', 'sheet_name': '2022', 'found_in_document': False},
        ],
        'script_versao': versao,
    }
    r = claude_svc.conferir(firm_id, user_id, _resultado(), 'revisao', planilha_beneficios=checagem)
    check('com a planilha: rascunho', bool(r.get('rascunho_id')), str(r)[:300])
    salvo = json.loads(FapReviewDraft.query.get(r['rascunho_id']).result_json)
    bloco = salvo.get('benefits_spreadsheet_review') or {}
    check('vai para o resultado como na tela (benefits_spreadsheet_review)', len(bloco.get('items', [])) == 2)
    check('contagens recontadas dos itens',
          (bloco.get('total_rows_with_theses'), bloco.get('found_count'), bloco.get('missing_count')) == (2, 1, 1),
          str(bloco)[:200])
    check('a resposta mostra o resumo da planilha',
          (r.get('planilha_beneficios') or {}).get('faltando') == 1, str(r.get('planilha_beneficios')))

    r = claude_svc.conferir(firm_id, user_id, json.dumps(_resultado()), 'revisao',
                            planilha_beneficios=json.dumps(checagem))
    check('planilha também como texto JSON', bool(r.get('rascunho_id')) and not r.get('aviso_planilha'),
          str(r.get('aviso_planilha')))

    feita_a_mao = dict(checagem)
    feita_a_mao.pop('script_versao')
    r = claude_svc.conferir(firm_id, user_id, _resultado(), 'revisao', planilha_beneficios=feita_a_mao)
    salvo = json.loads(FapReviewDraft.query.get(r['rascunho_id']).result_json)
    check('sem a versão do script: revisão segue, planilha fica fora, com aviso',
          'benefits_spreadsheet_review' not in salvo and 'script' in (r.get('aviso_planilha') or ''),
          str(r.get('aviso_planilha')))

    r = claude_svc.conferir(firm_id, user_id, _resultado(), 'revisao',
                            planilha_beneficios={'script_versao': versao, 'items': 'nada'})
    check('itens fora do formato: aviso, sem planilha', bool(r.get('aviso_planilha')))

    from mcp_server.tools.petition_reviewer import confer_review_handler, spreadsheet_script_handler
    r = confer_review_handler(firm_id, user_id, _resultado(), 'revisao', None, None, checagem)
    check('MCP: a planilha passa pelo handler da conferência',
          (r.get('planilha_beneficios') or {}).get('total') == 2, str(r.get('planilha_beneficios')))
    check('MCP: handler do script devolve a versão atual', spreadsheet_script_handler()['versao'] == versao)
    como_rodar = spreadsheet_script_handler()['como_rodar']
    check('como rodar diz onde ficam os anexos (uploads e attachments) e manda procurar',
          all(t in como_rodar for t in ('/mnt/user-data/uploads', '/mnt/user-data/attachments', 'find ')),
          como_rodar)

    r = claude_svc.conferir(firm_id, user_id, _resultado(), 'revisao')
    check('sem planilha: nada muda', 'aviso_planilha' not in r and 'planilha_beneficios' not in r)
    check('modelo de apresentação tem a seção da planilha', '📊' in r.get('como_mostrar', ''))


def test_descricoes():
    print('\n3) Instruções e descrições')
    instr = claude_svc._ADAPTACAO + claude_svc._ETAPAS['revisao'] + claude_svc._ETAPAS['comparacao']
    check('etapa de leitura manda rodar a checagem quando houver planilha',
          'conferir_planilha_beneficios' in claude_svc._ETAPAS['revisao']
          and 'conferir_planilha_beneficios' in claude_svc._ETAPAS['comparacao'], instr[:0])
    s = claude_svc.script_planilha()
    check('ferramenta explica como rodar e o que fazer sem ambiente de código',
          'python' in s['como_rodar'].lower() and 'não' in s['sem_ambiente_de_codigo'].lower())


def main():
    test_modulo_e_script()
    test_descricoes()
    with app.app_context():
        firm = LawFirm.query.order_by(LawFirm.id).first()
        user = User.query.filter(User.law_firm_id == firm.id).order_by(User.id).first()
        try:
            test_conferencia(firm.id, user.id)
        finally:
            db.session.rollback()
            FapReviewDraft.query.filter_by(law_firm_id=firm.id, user_id=user.id,
                                           saved_execution_id=None).filter(
                FapReviewDraft.petition_id.is_(None)).delete(synchronize_session=False)
            db.session.commit()
    print('\n' + '=' * 60)
    if _falhas:
        print(f'❌ {len(_falhas)} verificação(ões) falharam:')
        for f in _falhas:
            print(f'   · {f}')
        return 1
    print('✅ Todas as verificações passaram.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
