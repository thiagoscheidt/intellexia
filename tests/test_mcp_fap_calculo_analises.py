#!/usr/bin/env python3
"""
MCP do cálculo FAP — filtros novos da lista e do resumo, ranking, cruzamento
entre vigências, detalhe do estabelecimento e exportação para Excel.
Escritório descartável, criado e removido no fim.

    uv run python tests/test_mcp_fap_calculo_analises.py
"""

import os
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastmcp.exceptions import ToolError

from main import app
from app.models import (
    db, Benefit, FapCalculoEstabelecimento, FapCalculoItem, FapCompany, FapCompanyGroup,
    FapVigenciaCnpj, LawFirm, SeguradoNome,
)
from app.services import fap_group_service
from mcp_server.tools import exports
from mcp_server.tools.fap_calculo import (
    cruzamento_vigencias_calculo_fap_handler as cruzamento,
    detalhar_estabelecimento_calculo_fap_handler as detalhar,
    listar_beneficios_calculo_fap_handler as listar,
    ranking_calculo_fap_handler as ranking,
    resumo_beneficios_calculo_fap_handler as resumo,
)

TEST_CNPJ = '00000000001670'
RAIZ = '11222333'
EST1 = '11222333000181'
EST2 = '11222333000262'
RAIZ_B = '44555666'
EST_B = '44555666000110'

_falhas = []


def check(nome: str, condicao: bool, detalhe: str = '') -> None:
    if condicao:
        print(f'  ✅ {nome}')
    else:
        print(f'  ❌ {nome}{" — " + detalhe if detalhe else ""}')
        _falhas.append(nome)


def _item(firm, ano, tipo, numero, cnpj=EST1, **kw):
    base = dict(law_firm_id=firm.id, ano_vigencia=ano, cnpj=cnpj, cnpj_raiz=cnpj[:8], tipo=tipo,
                identificador=numero, estado_atual='ORIGINAL')
    base.update(kw)
    return FapCalculoItem(**base)


def setup():
    firm = LawFirm.query.filter_by(cnpj=TEST_CNPJ).first()
    if firm:
        cleanup(firm)
    firm = LawFirm(name='Escritório Teste MCP Cálculo Análises', cnpj=TEST_CNPJ)
    db.session.add(firm)
    db.session.flush()
    vig24 = FapVigenciaCnpj(law_firm_id=firm.id, employer_cnpj=EST1, vigencia_year='2024')
    db.session.add(vig24)
    db.session.flush()
    db.session.add_all([
        FapCompany(law_firm_id=firm.id, cnpj=RAIZ, nome='Empresa Calc Teste SA', synced_at=datetime.now()),
        FapCompany(law_firm_id=firm.id, cnpj=RAIZ_B, nome='Outra Calc Teste SA', synced_at=datetime.now()),
        # contestado em 2024 (Central), não em 2023
        Benefit(law_firm_id=firm.id, benefit_number='6444444444', benefit_type='B91',
                fap_vigencia_cnpj_id=vig24.id),
        # 2023
        _item(firm, 2023, 'B91', '6111111111', cpf='01234567890', total_pago=Decimal('100.00'),
              data_inicio=date(2021, 3, 1), cnpj_original=EST1, cnpj_atribuido=EST1),
        _item(firm, 2023, 'B94', '6222222222', estado_atual='EXCLUIDO', total_pago=Decimal('40.00')),
        _item(firm, 2023, 'B91', '6444444444', total_pago=Decimal('70.00')),
        _item(firm, 2023, 'CAT', '9990001'),
        # 2024
        _item(firm, 2024, 'B94', '6222222222', total_pago=Decimal('40.00'), data_inicio=date(2022, 6, 1)),
        _item(firm, 2024, 'B91', '6333333333', cnpj=EST2, total_pago=Decimal('900.00'),
              cnpj_original=EST1, cnpj_atribuido=EST2, data_obito=date(2024, 5, 1),
              data_cessacao=date(2024, 5, 1), calculo_id=400),
        _item(firm, 2024, 'B91', '6444444444', total_pago=Decimal('70.00')),
        _item(firm, 2024, 'B92', '6555555555', cnpj=EST_B, total_pago=Decimal('500.00')),
        FapCalculoEstabelecimento(
            law_firm_id=firm.id, ano_vigencia=2024, cnpj=EST2, cnpj_raiz=RAIZ, calculo_atual_id=500,
            calculo_original_id=400, calculo_atual_tipo='RECALCULO_ELETRONICO_PRIMEIRA_INSTANCIA',
            calculo_atual_data=date(2024, 10, 1), protocolo='PROT-777', situacao='ok',
            historico_json='[{"id": 500, "tipo": "RECALCULO_ELETRONICO_PRIMEIRA_INSTANCIA", '
                           '"data": "2024-10-01"}, {"id": 400, "tipo": "CALCULO_ORIGINAL", "data": "2024-09-01"}]'),
        SeguradoNome(law_firm_id=firm.id, cpf='01234567890', nome='MARIA DE TESTE',
                     status=SeguradoNome.STATUS_ENCONTRADO, fonte='teste', consultado_em=datetime.utcnow()),
    ])
    db.session.commit()
    fap_group_service.assign_group(firm.id, RAIZ, 'Grupo Calc Teste')
    db.session.commit()
    return firm


def cleanup(firm):
    for modelo in (FapCalculoItem, FapCalculoEstabelecimento, Benefit, FapVigenciaCnpj, FapCompany,
                   FapCompanyGroup, SeguradoNome):
        modelo.query.filter_by(law_firm_id=firm.id).delete()
    db.session.delete(firm)
    db.session.commit()


def _ids(r):
    return sorted(f"{i['ano_vigencia']}:{i['identificador']}" for i in r['itens'])


def test_filtros(f):
    print('\n1. filtros novos da lista')
    r = listar(f.id, tipo='B91,B94')
    check('tipo com vírgula', r['total_encontrado'] == 6, str(r['total_encontrado']))
    r = listar(f.id, tipo=['b92', 'CAT'])
    check('tipo em lista (e minúsculo)', r['total_encontrado'] == 2, str(r['total_encontrado']))
    try:
        listar(f.id, tipo='B91,B99')
        check('tipo inválido na lista é erro', False)
    except ToolError as e:
        check('tipo inválido na lista é erro e cita o inválido', 'B99' in str(e), str(e))

    r = listar(f.id, ordem='total_pago')
    valores = [i['total_pago'] for i in r['itens']]
    check('ordem total_pago: maior primeiro, vazio no fim', valores[0] == 900.0 and valores[-1] is None,
          str(valores))
    try:
        listar(f.id, ordem='aleatoria')
        check('ordem inválida é erro', False)
    except ToolError as e:
        check('ordem inválida lista as aceitas', 'total_pago' in str(e), str(e))

    r = listar(f.id, cnpj_divergente=True)
    check('cnpj_divergente', _ids(r) == ['2024:6333333333'], str(_ids(r)))
    r = listar(f.id, cnpj_divergente=False, vigencia=2024)
    check('cnpj_divergente=false exclui só o divergente', r['total_encontrado'] == 3, str(_ids(r)))

    r = listar(f.id, ausente_no_calculo_atual=True)
    check('ausente_no_calculo_atual', _ids(r) == ['2024:6333333333'], str(_ids(r)))
    check('item ausente marcado', r['itens'][0]['ausente_no_calculo_atual'] is True)

    r = listar(f.id, inicio_de='2022-01-01')
    check('início a partir de', _ids(r) == ['2024:6222222222'], str(_ids(r)))
    r = listar(f.id, inicio_de='01/01/2021', inicio_ate='2021-12-31')
    check('início entre (DD/MM/AAAA e ISO)', _ids(r) == ['2023:6111111111'], str(_ids(r)))
    r = listar(f.id, cessacao_ate='2024-12-31')
    check('cessação até', _ids(r) == ['2024:6333333333'], str(_ids(r)))
    try:
        listar(f.id, inicio_de='ontem')
        check('data inválida é erro', False)
    except ToolError as e:
        check('data inválida explica o formato', 'AAAA-MM-DD' in str(e), str(e))
    r = listar(f.id, com_obito=True)
    check('com óbito', _ids(r) == ['2024:6333333333'], str(_ids(r)))

    r = listar(f.id, nome='maria')
    check('nome revelado acha o CPF', _ids(r) == ['2023:6111111111'], str(_ids(r)))
    r = listar(f.id, nome='ninguém com esse nome')
    check('nome sem revelação: vazio, não tudo', r['total_encontrado'] == 0)


def test_resumo(f):
    print('\n2. resumo com os filtros da lista')
    s = resumo(f.id)
    check('total', s['total'] == 8, str(s['total']))
    check('conta + excluídos', s['conta'] == 7 and s['excluidos'] == 1, f"{s['conta']} {s['excluidos']}")
    check('contestados (por vigência)', s['contestados'] == 1, str(s['contestados']))
    check('conta e não contestados', s['conta_nao_contestados'] == 6, str(s['conta_nao_contestados']))
    check('total pago dos que contam sem contestação',
          s['total_pago_conta_nao_contestado'] == Decimal('1610.00'), str(s['total_pago_conta_nao_contestado']))

    s = resumo(f.id, situacao='conta', contestado=False)
    check('situacao + contestado no resumo', s['total'] == 6 and s['total_pago'] == Decimal('1610.00'),
          f"{s['total']} {s['total_pago']}")
    for kw in ({'situacao': 'excluido'}, {'contestado': True}, {'tipo': 'B91,B94'}, {'nome': 'maria'},
               {'cnpj_divergente': True}, {'ausente_no_calculo_atual': True}, {'com_obito': False},
               {'inicio_de': '2022-01-01'}, {'cpf': '012.345.678-90'}, {'numero': '6444444444'}):
        check(f'resumo bate com a lista {kw}',
              resumo(f.id, **kw)['total'] == listar(f.id, **kw)['total_encontrado'])


def test_ranking(f):
    print('\n3. ranking')
    r = ranking(f.id)
    check('padrão: por empresa, pelo valor que conta sem contestação',
          r['agrupado_por'] == 'empresa' and r['ordenado_por'] == 'total_pago_conta_nao_contestado')
    check('duas empresas, a de maior valor primeiro',
          [b['cnpj_raiz'] for b in r['itens']] == [RAIZ, RAIZ_B], str([b['cnpj_raiz'] for b in r['itens']]))
    check('nome da empresa', r['itens'][0]['empresa_nome'] == 'Empresa Calc Teste SA')
    check('valores em texto, como no resumo', r['itens'][0]['total_pago_conta_nao_contestado'] == '1110.00',
          str(r['itens'][0]['total_pago_conta_nao_contestado']))
    total = resumo(f.id)['total']
    check('soma dos blocos = resumo', sum(b['total'] for b in r['itens']) == total)

    r = ranking(f.id, agrupar_por='estabelecimento', empresa='Empresa Calc', ordenar_por='total')
    check('por estabelecimento dentro da empresa', [b['cnpj'] for b in r['itens']] == [EST1, EST2],
          str([(b['cnpj'], b['total']) for b in r['itens']]))

    r = ranking(f.id, agrupar_por='vigencia', ordenar_por='excluidos')
    check('por vigência', [b['ano_vigencia'] for b in r['itens']] == [2023, 2024],
          str([(b['ano_vigencia'], b['excluidos']) for b in r['itens']]))

    r = ranking(f.id, agrupar_por='grupo', ordenar_por='total')
    nomes = [b['grupo_nome'] for b in r['itens']]
    check('por grupo, com "Sem grupo"', nomes == ['Grupo Calc Teste', 'Sem grupo'], str(nomes))

    r = ranking(f.id, agrupar_por='empresa', limite=1, deslocamento=1)
    check('paginação de blocos', r['total_encontrado'] == 2 and r['retornados'] == 1
          and r['itens'][0]['cnpj_raiz'] == RAIZ_B and not r['tem_mais'])
    for kw in ({'agrupar_por': 'cidade'}, {'ordenar_por': 'nome'}):
        try:
            ranking(f.id, **kw)
            check(f'ToolError {kw}', False)
        except ToolError as e:
            check(f'ToolError {kw} lista aceitos', 'aceitos' in str(e), str(e))


def test_cruzamento(f):
    print('\n4. cruzamento entre vigências')
    r = cruzamento(f.id)
    check('excluído em 2023, contando sem contestação em 2024', _ids(r) == ['2024:6222222222'], str(_ids(r)))
    outras = r['itens'][0]['outras_vigencias']
    check('outras_vigencias traz a de 2023 excluída',
          [(v['ano_vigencia'], v['situacao']) for v in outras] == [(2023, 'excluido')], str(outras))
    check('total pago do recorte', r['total_pago'] == '40.00', r['total_pago'])

    r = cruzamento(f.id, criterio='contestado')
    check('contestado em 2024, não contestado em 2023', _ids(r) == ['2023:6444444444'], str(_ids(r)))
    check('a outra vigência aparece contestada',
          r['itens'][0]['outras_vigencias'][0]['contestado'] is True, str(r['itens'][0]['outras_vigencias']))

    r = cruzamento(f.id, criterio='qualquer')
    check('qualquer = os dois', _ids(r) == ['2023:6444444444', '2024:6222222222'], str(_ids(r)))
    r = cruzamento(f.id, vigencia=2023)
    check('filtro de vigência é a do alvo', r['total_encontrado'] == 0)
    try:
        cruzamento(f.id, criterio='talvez')
        check('critério inválido é erro', False)
    except ToolError as e:
        check('critério inválido lista aceitos', 'excluido' in str(e), str(e))


def test_estabelecimento(f):
    print('\n5. detalhe do estabelecimento')
    d = detalhar(f.id, '11.222.333/0002-62')
    check('vigência mais recente', d['ano_vigencia'] == 2024 and d['vigencias_disponiveis'] == [2024])
    check('cálculo atual e recalculado', d['calculo_atual']['rotulo'] == 'Recálculo 1ª instância'
          and d['recalculado'] is True and d['protocolo'] == 'PROT-777', str(d['calculo_atual']))
    check('histórico', [h['id'] for h in d['historico']] == [500, 400])
    check('contadores em texto', d['contadores']['total'] == 1 and d['contadores']['total_pago'] == '900.00')
    check('itens ausentes do cálculo atual', d['ausentes_no_calculo_atual']['total'] == 1
          and d['ausentes_no_calculo_atual']['itens'][0]['identificador'] == '6333333333')
    check('nenhum excluído', d['excluidos']['total'] == 0)
    check('dica aponta a lista', 'listar_beneficios_calculo_fap' in d['dica'])

    d = detalhar(f.id, EST1, 2023)
    check('estabelecimento sem registro de sincronização ainda traz os itens',
          d['contadores']['total'] == 4 and d['calculo_atual'] is None and d['excluidos']['total'] == 1)
    check('vigências do estabelecimento', d['vigencias_disponiveis'] == [2024, 2023])

    check('CNPJ desconhecido', 'não encontrado' in detalhar(f.id, '99888777000166').get('erro', ''))
    r = detalhar(f.id, EST1, 2019)
    check('vigência sem cálculo lista as que existem', '2024, 2023' in r.get('erro', ''), str(r))
    try:
        detalhar(f.id, '11222333')
        check('raiz no lugar do CNPJ é erro', False)
    except ToolError as e:
        check('raiz no lugar do CNPJ aponta o ranking', 'ranking_calculo_fap' in str(e), str(e))


def test_exportar(f):
    print('\n6. exportar para Excel')
    r = exports.export_calculo_excel_handler(f.id, 'https://mcp.exemplo', vigencia=2024, tipo='B91',
                                             ordem='total_pago')
    check('link e linhas', r.get('total_linhas') == 2 and '/export/' in r.get('url_download', ''), str(r))
    caminho = os.path.join(exports._EXPORT_DIR, str(f.id), r['arquivo'])
    check('arquivo gravado', os.path.isfile(caminho))
    from openpyxl import load_workbook
    ws = load_workbook(caminho, read_only=True).active
    linhas = [list(x) for x in ws.iter_rows(values_only=True)]
    filtros = ' '.join(str(c) for linha in linhas[:4] for c in linha if c)
    check('filtros no cabeçalho', 'Vigência:' in filtros and '2024' in filtros and 'B91' in filtros, filtros)
    nbs = [linha[4] for linha in linhas if linha and linha[4] in ('6333333333', '6444444444')]
    check('ordem pedida (maior total pago primeiro)', nbs == ['6333333333', '6444444444'], str(nbs))
    os.remove(caminho)
    r = exports.export_calculo_excel_handler(f.id, 'https://mcp.exemplo', vigencia=2019)
    check('sem itens: erro, sem arquivo', r.get('total_linhas') == 0 and 'erro' in r)


def main() -> int:
    with app.app_context():
        firm = setup()
        try:
            test_filtros(firm)
            test_resumo(firm)
            test_ranking(firm)
            test_cruzamento(firm)
            test_estabelecimento(firm)
            test_exportar(firm)
        finally:
            db.session.rollback()
            cleanup(firm)

    print('\n' + '=' * 62)
    if _falhas:
        print(f'❌ {len(_falhas)} verificação(ões) falharam:')
        for n in _falhas:
            print(f'   - {n}')
        return 1
    print('✅ Tudo verde')
    return 0


if __name__ == '__main__':
    sys.exit(main())
