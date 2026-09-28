#!/usr/bin/env python3
"""
Ferramentas MCP da Base de Jurisprudência: pesquisa (campos e inteiro teor),
detalhe, panorama, valores de filtro, decisões parecidas e exportação.

Roda contra SQLite descartável; os índices (inteiro teor e parecidas) são
simulados — o teste confere o contrato da ferramenta, não o Meilisearch/Qdrant.

Uso: uv run python tests/test_mcp_jurisprudencia.py
"""
import os
import sys
import tempfile
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from main import app

DB_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_tmp_mcp_jurisprudencia.db')
if os.path.exists(DB_FILE):
    os.remove(DB_FILE)
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{DB_FILE}'

from app.models import db  # noqa: E402

app.extensions.pop('sqlalchemy', None)
try:
    db._app_engines.pop(app, None)
except Exception:
    pass
db.init_app(app)

from fastmcp.exceptions import ToolError  # noqa: E402

from app.services import jurisprudence_index_service as indice  # noqa: E402
from mcp_server.tools import exports  # noqa: E402
from mcp_server.tools.jurisprudence import (  # noqa: E402
    decision_detail_handler,
    export_jurisprudence_excel_handler,
    jurisprudence_filter_values_handler,
    jurisprudence_overview_handler,
    search_jurisprudence_handler,
    similar_decisions_handler,
)

PASTA_EXPORT = tempfile.mkdtemp(prefix='mcp_jur_export_')
exports._EXPORT_DIR = PASTA_EXPORT
APP_URL = 'https://app.exemplo'
FALHAS = []


def check(rotulo, condicao, extra=''):
    print(f"  [{'OK ' if condicao else 'FALHA'}] {rotulo}{(' — ' + str(extra)) if extra else ''}")
    if not condicao:
        FALHAS.append(rotulo)


def erro(funcao, *args, **kwargs):
    try:
        funcao(*args, **kwargs)
    except ToolError as e:
        return str(e)
    return None


def _preparar():
    from app.models import LawFirm, JudicialLegalThesis
    from app.services import jurisprudence_service as svc
    from app.services import jurisprudence_thesis_service as ts

    db.create_all()
    db.session.add_all([LawFirm(id=1, name='E', cnpj='00000000000191'), LawFirm(id=2, name='O', cnpj='00000000000272')])
    db.session.add(JudicialLegalThesis(id=1, law_firm_id=1, key='trajeto_b91', name='TRAJETO - B91', is_active=True))
    db.session.commit()
    r = svc.ResolvedorDeTeses(1)

    def decisao(processo, tipo, resultado, data, teses, tribunal='TRF4', orgao='2ª Turma', motivo=''):
        svc.criar_decisao(1, {'processo': processo, 'tipo_documento': tipo, 'resultado': resultado,
                              'data_julgamento': data, 'teses': teses, 'tribunal': tribunal,
                              'orgao_julgador': orgao, 'motivo_resultado': motivo,
                              'classe_processual': 'Apelação Cível', 'vigencia_fap': '2017 a 2021',
                              'precedentes': ['TRF4 AC 5003321-10.2023.4.04.7207']},
                          source='planilha', resolvedor=r)

    # virada: sentença desfavorável, acórdão favorável
    decisao('5012345-67.2021.4.04.7205', 'SENTENCA', 'DESFAVORAVEL', '2024-03-14', ['Acidente de Trajeto'],
            orgao='2ª Vara Federal de Blumenau', motivo='o percurso casa-trabalho integra o risco')
    decisao('5012345-67.2021.4.04.7205', 'ACORDAO', 'FAVORAVEL', '2026-05-12',
            ['ACIDENTE DE TRAJETO', 'PRORROGAÇÃO DE BENEFÍCIO'], motivo='excluído o acidente de trajeto')
    decisao('5003321-10.2023.4.04.7207', 'SENTENCA', 'PARCIALMENTE FAVORAVEL', '2025-02-02',
            ['ACIDENTE DE TRAJETO', 'PRORROGACAO DE BENEFICIO'], orgao='1ª Vara Federal de Tubarão',
            motivo='exclusão dos eventos de trajeto')
    decisao('5006749-11.2023.4.03.6114', 'EMBARGOS DE DECLARACAO', 'DESFAVORAVEL', '2024-10-10',
            ['BIS IN IDEM'], tribunal='TRF3', orgao='7ª Vara Cível Federal de São Paulo')
    db.session.commit()
    from app.models import JurisprudenceThesis as T
    ts.ligar(1, T.query.filter_by(law_firm_id=1, key='ACIDENTE DE TRAJETO').first().id, [1])


def test_pesquisa():
    print('\n1. pesquisar_jurisprudencia')
    tudo = search_jurisprudence_handler(1, APP_URL)
    check('sem termo lista a base inteira', tudo['total_encontrado'] == 4 and tudo['processos_distintos'] == 3)
    check('envelope paginado padrão', {'total_encontrado', 'retornados', 'tem_mais', 'itens'} <= set(tudo))
    item = tudo['itens'][0]
    check('item traz citação, teses nos dois campos e link',
          item['citacao'].startswith('(') and 'teses_originais' in item and 'teses_catalogo' in item
          and item['url'].startswith(APP_URL + '/process-panel/jurisprudencia/decisao/'), item)
    p1 = search_jurisprudence_handler(1, APP_URL, limite=2)
    p2 = search_jurisprudence_handler(1, APP_URL, limite=2, deslocamento=p1['proximo_deslocamento'])
    ids = [i['id'] for i in p1['itens'] + p2['itens']]
    check('paginação cobre tudo sem repetir', sorted(ids) == sorted(i['id'] for i in tudo['itens']) and len(set(ids)) == 4)
    check('sem próxima página no fim', p2['tem_mais'] is False)

    r = search_jurisprudence_handler(1, APP_URL, termo='percurso')
    check('sinônimo: "percurso" acha trajeto', r['total_encontrado'] == 3, r['total_encontrado'])
    check('trecho vem em texto, com «» no achado', any('«' in (i['trecho'] or '') for i in r['itens']),
          [i['trecho'] for i in r['itens']])
    check('número sem pontuação acha o processo',
          search_jurisprudence_handler(1, APP_URL, termo='50123456720214047205')['total_encontrado'] == 2)
    check('como_tem_decidido soma o recorte',
          sum(v['decisoes'] for v in r['como_tem_decidido'].values()) == r['total_encontrado'])

    check('tese do catálogo pelo nome, sem acento e em minúsculas',
          search_jurisprudence_handler(1, APP_URL, teses_catalogo=['trajeto - b91'])['total_encontrado'] == 3)
    msg = erro(search_jurisprudence_handler, 1, APP_URL, teses_catalogo=['Tese Inventada'])
    check('tese do catálogo inexistente → erro que manda consultar os valores', msg and 'valores_de_filtro' in msg, msg)
    check('tese original junta PRORROGAÇÃO e PRORROGACAO',
          search_jurisprudence_handler(1, APP_URL, teses_originais=['prorrogação de benefício'])['total_encontrado'] == 2)
    check('resultado aceita "favorável" com acento',
          search_jurisprudence_handler(1, APP_URL, resultados=['favorável'])['total_encontrado'] == 1)
    check('instância e tribunal',
          search_jurisprudence_handler(1, APP_URL, instancias=['sentença'], tribunais=['trf4'])['total_encontrado'] == 2)
    check('vigência coberta pelo intervalo',
          search_jurisprudence_handler(1, APP_URL, vigencia=2019)['total_encontrado'] == 4)
    check('data inválida → erro com o formato', 'YYYY-MM-DD' in (erro(search_jurisprudence_handler, 1, APP_URL,
                                                                        julgada_de='ontem') or ''))
    check('modo inválido → erro', erro(search_jurisprudence_handler, 1, APP_URL, modo='tudo') is not None)
    check('outro escritório vê a base vazia', search_jurisprudence_handler(2, APP_URL)['total_encontrado'] == 0)


def test_inteiro_teor():
    print('\n2. Inteiro teor (índice simulado)')
    from app.models import JurisprudenceDecision as D
    alvo = D.query.filter_by(law_firm_id=1, tipo_documento='acordao').first()
    original = indice.buscar_inteiro_teor
    indice.buscar_inteiro_teor = lambda *a, **k: [{'decision_id': alvo.id, 'kind': 'pagina', 'page': 3,
                                                   'trecho': 'afastado o <mark>trajeto</mark> &amp; mais'}]
    try:
        r = search_jurisprudence_handler(1, APP_URL, termo='trajeto', modo='inteiro_teor', uf='SC')
        check('devolve a decisão com página e trecho em texto',
              r['itens'][0]['pagina'] == 3 and r['itens'][0]['trecho'] == 'afastado o «trajeto» & mais', r['itens'])
        check('avisa quantas decisões têm inteiro teor', 'cobertura_inteiro_teor' in r)
        check('avisa o filtro que não vale no inteiro teor', 'uf' in (r.get('filtros_ignorados') or ''))
        check('sem termo → erro', erro(search_jurisprudence_handler, 1, APP_URL, modo='inteiro_teor') is not None)
        indice.buscar_inteiro_teor = lambda *a, **k: None
        check('índice fora → erro que sugere o modo campos',
              'campos' in (erro(search_jurisprudence_handler, 1, APP_URL, termo='x', modo='inteiro_teor') or ''))
    finally:
        indice.buscar_inteiro_teor = original


def test_detalhe_e_panorama():
    print('\n3. detalhar_decisao e panorama_jurisprudencia')
    from app.models import JurisprudenceDecision as D
    acordao = D.query.filter_by(law_firm_id=1, tipo_documento='acordao').first()
    d = decision_detail_handler(acordao.id, 1, APP_URL)
    check('detalhe traz fundamentos, argumentos e precedentes',
          {'fundamentos', 'argumentos_acolhidos', 'argumentos_rejeitados', 'precedentes_citados'} <= set(d))
    check('precedente que está na base vem com o id',
          d['precedentes_citados'][0]['decisao_na_base_id'] is not None, d['precedentes_citados'])
    check('traz a sentença do mesmo processo', [x['instancia'] for x in d['outras_decisoes_do_processo']] == ['Sentença'])
    check('tese original como veio', d['teses_originais'] == ['ACIDENTE DE TRAJETO', 'PRORROGAÇÃO DE BENEFÍCIO'])
    check('decisão de outro escritório → erro', erro(decision_detail_handler, acordao.id, 2) is not None)

    p = jurisprudence_overview_handler(1, APP_URL, teses_originais=['acidente de trajeto'])
    check('panorama conta o recorte', p['decisoes'] == 3 and p['processos'] == 2, (p['decisoes'], p['processos']))
    check('panorama por tribunal', p['por_tribunal'][0]['grupo'] == 'TRF4' and p['por_tribunal'][0]['decisoes'] == 3)
    check('panorama acha a virada no acórdão', p['total_viradas'] == 1 and p['viradas_no_acordao'][0]['acordao_id'] == acordao.id)
    check('favoráveis mais recentes, sem desfavorável',
          [i['resultado'] for i in p['favoraveis_mais_recentes']] == ['Favorável', 'Parcialmente favorável'])


def test_valores_parecidas_exportacao():
    print('\n4. valores de filtro, parecidas e exportação')
    v = jurisprudence_filter_values_handler(1)
    check('valores trazem os dois campos de tese',
          v['teses_catalogo'][0] == {'nome': 'TRAJETO - B91', 'decisoes': 3}
          and any(t['nome'] == 'ACIDENTE DE TRAJETO' and t['decisoes'] == 3 for t in v['teses_originais']), v['teses_catalogo'])
    check('período de julgamento da base', v['periodo_de_julgamento'] == {'de': '2024-03-14', 'ate': '2026-05-12'})

    from app.models import JurisprudenceDecision as D
    a, b = D.query.filter_by(law_firm_id=1).order_by(D.id).limit(2).all()
    original = indice.parecidas
    indice.parecidas = lambda d, limite=6: [{'decision_id': b.id, 'score': 0.87654}]
    try:
        r = similar_decisions_handler(a.id, 1, APP_URL)
        check('parecidas devolve a decisão com a semelhança', r['parecidas'][0]['id'] == b.id
              and r['parecidas'][0]['semelhanca'] == 0.877)
        indice.parecidas = lambda d, limite=6: None
        check('fora do índice → erro que sugere pesquisar', 'pesquisar_jurisprudencia' in
              (erro(similar_decisions_handler, a.id, 1) or ''))
    finally:
        indice.parecidas = original

    from openpyxl import load_workbook
    r = export_jurisprudence_excel_handler(1, 'https://mcp.exemplo', APP_URL, tribunais=['TRF4'])
    check('exporta com link assinado', r['total_linhas'] == 3 and r['url_download'].startswith('https://mcp.exemplo/export/'))
    caminho = next(os.path.join(raiz, n) for raiz, _, nomes in os.walk(PASTA_EXPORT) for n in nomes)
    ws = load_workbook(caminho).active
    cabecalho = [c.value for c in ws[1]]
    check('planilha com os dois campos de tese e a citação',
          {'Teses originais', 'Teses do catálogo', 'Citação'} <= set(cabecalho) and ws.max_row == 4)
    check('exportação vazia avisa', export_jurisprudence_excel_handler(2, 'https://mcp.exemplo')['total_linhas'] == 0)


def test_registro():
    print('\n5. Registro no servidor MCP')
    import asyncio
    from mcp_server import server
    nomes = {t.name for t in asyncio.run(server.mcp.list_tools())}
    esperadas = {'pesquisar_jurisprudencia', 'detalhar_decisao', 'panorama_jurisprudencia',
                 'valores_de_filtro_jurisprudencia', 'decisoes_parecidas', 'exportar_jurisprudencia_excel'}
    check('as seis ferramentas estão registradas', esperadas <= nomes, esperadas - nomes)


def main():
    print('=' * 60)
    print('MCP · BASE DE JURISPRUDÊNCIA')
    print('=' * 60)
    with app.app_context():
        assert DB_FILE in str(db.engine.url), 'ABORTADO: fora do sandbox'
        _preparar()
        test_pesquisa()
        test_inteiro_teor()
        test_detalhe_e_panorama()
        test_valores_parecidas_exportacao()
    test_registro()
    print('\n' + '=' * 60)
    if FALHAS:
        print(f'❌ {len(FALHAS)} falha(s): {", ".join(FALHAS)}')
        return 1
    print('✅ Todos os testes passaram')
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    finally:
        import shutil
        if os.path.exists(DB_FILE):
            os.remove(DB_FILE)
        shutil.rmtree(PASTA_EXPORT, ignore_errors=True)
